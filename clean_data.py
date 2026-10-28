"""
بننضّف بيانات العملاء وبنطلّع ملف Excel نهائي (بيانات تجريبية).
الملف اللي داخل : dirty_customers.csv
الملفات اللي طالعة: customers_cleaned.xlsx  (شيت البيانات النضيفة + تقرير التنضيف + ملخص برسوم)
         cleaned_customers.csv
"""
import re
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = "dirty_customers.csv"
OUT = "customers_cleaned.xlsx"

df = pd.read_csv(SRC, dtype=str, keep_default_na=True)
rows_before = len(df)
report = []  # هنسجّل هنا كل مشكلة: (المشكلة, عدد الحالات, الحل)


def log(problem, count, fix):
    report.append((problem, int(count), fix))


# ---------- 0) بنبص على الداتا الأول ----------
missing_before = df.isna().sum()

# ---------- 1) الأسماء: مسافات زايدة ----------
orig = df["name"].copy()
df["name"] = df["name"].str.strip().str.replace(r"\s+", " ", regex=True)
log("مسافات زائدة في الأسماء", (orig != df["name"]).sum(), "strip + توحيد المسافات بـ regex")

# ---------- 2) المدن: غلطات إملاء وأسماء بالإنجليزي ----------
city_map = {
    "القاهرة": ["القاهره", "cairo", "القاهرا"],
    "الجيزة": ["الجيزه", "giza"],
    "الإسكندرية": ["الاسكندرية", "اسكندريه", "الأسكندرية", "alexandria", "alex"],
    "المنصورة": ["المنصوره", "mansoura"],
    "طنطا": ["طنطه", "tanta"],
    "أسيوط": ["اسيوط", "assiut"],
    "الأقصر": ["الاقصر", "luxor"],
    "أسوان": ["اسوان", "aswan"],
    "بورسعيد": ["بور سعيد", "port said"],
    "السويس": ["suez"],
}
lookup = {}
for good, variants in city_map.items():
    lookup[good.lower()] = good
    for v in variants:
        lookup[v.lower()] = good
city_raw = df["city"].copy()
df["city"] = df["city"].str.strip().str.lower().map(lookup)
fixed_city = ((city_raw.notna()) & (city_raw != df["city"])).sum()
unmapped = (city_raw.notna() & df["city"].isna()).sum()
log("أخطاء إملائية/صيغ إنجليزية في أسماء المدن", fixed_city, "قاموس تصحيح (mapping) بعد توحيد الحروف والمسافات")
if unmapped:
    log("مدن غير معروفة", unmapped, "تحويلها لقيمة ناقصة")

# ---------- 3) الجنس ----------
gmap = {"ذكر": "ذكر", "male": "ذكر", "m": "ذكر", "أنثى": "أنثى", "انثى": "أنثى", "female": "أنثى", "f": "أنثى"}
g_raw = df["gender"].copy()
df["gender"] = df["gender"].str.strip().str.lower().map(gmap)
log("صيغ مختلفة لعمود الجنس (M/male/ذكر...)", ((g_raw.notna()) & (g_raw != df["gender"])).sum(), "توحيد إلى (ذكر / أنثى)")

# ---------- 4) الموبايل ----------
def clean_phone(p):
    if pd.isna(p):
        return np.nan
    digits = re.sub(r"\D", "", p)
    if digits.startswith("0020"):
        digits = digits[4:]
    elif digits.startswith("20") and len(digits) == 12:
        digits = digits[2:]
    if len(digits) == 10 and digits[0] == "1":
        digits = "0" + digits
    return digits if re.fullmatch(r"01[0125]\d{8}", digits) else np.nan

ph_raw = df["phone"].copy()
ph_new = ph_raw.apply(clean_phone)
invalid_phone = (ph_raw.notna() & ph_new.isna()).sum()
reformatted = (ph_raw.notna() & ph_new.notna() & (ph_raw != ph_new)).sum()
df["phone"] = ph_new
log("أرقام موبايل بصيغ مختلفة (+20 / 0020 / شرطات / مسافات)", reformatted, "إزالة غير الأرقام ثم توحيد الصيغة 01XXXXXXXXX")
log("أرقام موبايل غير صالحة", invalid_phone, "التحقق بـ regex وتحويل غير الصالح لقيمة ناقصة")

# ---------- 5) الإيميل ----------
em_raw = df["email"].copy()
df["email"] = df["email"].str.strip().str.lower()
log("إيميلات بحروف كبيرة", ((em_raw.notna()) & (em_raw != df["email"])).sum(), "تحويل لحروف صغيرة")
valid_email = df["email"].str.fullmatch(r"[a-z0-9._%+-]+@[a-z0-9.-]+\.[a-z]{2,}", na=False)
bad_email = (df["email"].notna() & ~valid_email).sum()
df.loc[~valid_email, "email"] = np.nan
log("إيميلات غير صالحة", bad_email, "التحقق بـ regex وتحويل غير الصالح لقيمة ناقصة")

# ---------- 6) التواريخ ----------
def parse_date(s):
    if pd.isna(s):
        return pd.NaT
    s = s.strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d-%m-%Y", "%d %b %Y"):
        try:
            return pd.to_datetime(s, format=fmt)
        except (ValueError, TypeError):
            continue
    return pd.NaT

d_raw = df["signup_date"].copy()
df["signup_date"] = d_raw.apply(parse_date)
log("تواريخ بصيغ متعددة", (d_raw.notna() & df["signup_date"].notna() & ~d_raw.str.fullmatch(r"\d{4}-\d{2}-\d{2}", na=False)).sum(),
    "تجربة عدة صيغ وتحويلها إلى تاريخ موحد YYYY-MM-DD (يوم/شهر/سنة للصيغ المصرية)")
log("تواريخ مستحيلة/غير صالحة (31/02، 0000-00-00...)", (d_raw.notna() & df["signup_date"].isna()).sum(), "تحويلها لقيمة ناقصة")

# ---------- 7) الأرقام والقيم الغريبة ----------
df["age"] = pd.to_numeric(df["age"], errors="coerce")
df["total_spent"] = pd.to_numeric(df["total_spent"], errors="coerce")
df["customer_id"] = pd.to_numeric(df["customer_id"], errors="coerce").astype(int)

bad_age = ((df["age"] < 16) | (df["age"] > 90)).sum()
df.loc[(df["age"] < 16) | (df["age"] > 90), "age"] = np.nan
log("أعمار شاذة (250 / 0 / سالب / 999)", bad_age, "اعتبار أي عمر خارج 16–90 قيمة خاطئة وتحويله لناقصة ثم تعويضه")

neg = (df["total_spent"] < 0).sum()
df.loc[df["total_spent"] < 0, "total_spent"] = np.nan
upper = 100_000  # حد منطقي: أي عميل إنفاقه فوق 100 ألف يبقى غلط في الإدخال (قاعدة شغل مش قاعدة إحصائية)
huge = (df["total_spent"] > upper).sum()
df.loc[df["total_spent"] > upper, "total_spent"] = np.nan
log("مبالغ إنفاق سالبة", neg, "تحويلها لقيمة ناقصة ثم تعويضها")
log(f"مبالغ إنفاق مستحيلة (أكبر من {upper:,})", huge, "قاعدة عمل (حد أقصى منطقي) ثم تعويضها بالوسيط. تم تجنب IQR لأنه كان هيحذف قيم إنفاق عالية لكنها سليمة")

# ---------- 8) التكرارات ----------
# العميل المتكرر: يا نفس الـ customer_id (نسخة طبق الأصل) يا نفس الإيميل والاسم (بس بصيغة مختلفة)
n0 = len(df)
df = df.drop_duplicates(subset=["customer_id"], keep="first")
exact_removed = n0 - len(df)
n1 = len(df)
has_email = df["email"].notna()
dup_mask = has_email & df.duplicated(subset=["email", "name"], keep="first")
df = df[~dup_mask]
near_removed = n1 - len(df)
log("صفوف مكررة تماماً (نفس customer_id)", exact_removed, "drop_duplicates على customer_id")
log("صفوف مكررة بصيغة مختلفة (نفس الاسم والإيميل بعد التنظيف)", near_removed, "اكتشافها بعد توحيد الصيغ ثم حذفها")

# ---------- 9) القيم الفاضية ----------
age_missing = df["age"].isna().sum()
df["age"] = df["age"].fillna(df["age"].median()).round().astype(int)
spent_missing = df["total_spent"].isna().sum()
df["total_spent"] = df["total_spent"].fillna(df["total_spent"].median()).round(2)
log("أعمار ناقصة (أصلية + الشاذة)", age_missing, f"تعويض بالوسيط ({int(df['age'].median())} سنة)")
log("مبالغ إنفاق ناقصة (أصلية + الشاذة)", spent_missing, "تعويض بالوسيط")

city_missing = df["city"].isna().sum()
gender_missing = df["gender"].isna().sum()
df["city"] = df["city"].fillna("غير محدد")
df["gender"] = df["gender"].fillna("غير محدد")
log("مدن ناقصة", city_missing, "وضع القيمة 'غير محدد' لأن التعويض بالوسيط غير منطقي للنصوص")
log("جنس ناقص", gender_missing, "وضع القيمة 'غير محدد'")
log("موبايل / إيميل / تاريخ ناقص", int(df["phone"].isna().sum() + df["email"].isna().sum() + df["signup_date"].isna().sum()),
    "تُركت فارغة عمداً (لا يصح اختراع أرقام أو إيميلات)")

df = df.sort_values("customer_id").reset_index(drop=True)
df.columns = ["رقم_العميل", "الاسم", "الموبايل", "الإيميل", "المدينة", "تاريخ_التسجيل", "العمر", "الجنس", "إجمالي_الإنفاق"]
rows_after = len(df)

df.to_csv("cleaned_customers.csv", index=False, encoding="utf-8-sig")

# ---------- بنجهّز جداول الملخص ----------
by_city = df[df["المدينة"] != "غير محدد"].groupby("المدينة").size().sort_values(ascending=False).reset_index(name="عدد العملاء")
by_gender = df.groupby("الجنس").size().reset_index(name="عدد العملاء")
df["_m"] = df["تاريخ_التسجيل"].dt.to_period("M").astype(str)
by_month = df[df["_m"] != "NaT"].groupby("_m").size().reset_index(name="عدد التسجيلات").rename(columns={"_m": "الشهر"})
spent_city = (df[df["المدينة"] != "غير محدد"].groupby("المدينة")["إجمالي_الإنفاق"].sum().round(0)
              .sort_values(ascending=False).reset_index(name="إجمالي الإنفاق"))
df = df.drop(columns="_m")

# قبل/بعد
ba = pd.DataFrame({
    "المؤشر": ["عدد الصفوف", "إجمالي القيم الناقصة (أصلاً)"],
    "قبل": [rows_before, int(missing_before.sum())],
    "بعد": [rows_after, int(df[["الموبايل", "الإيميل", "تاريخ_التسجيل"]].isna().sum().sum())],
})

# ---------- بنكتب ملف Excel ----------
rep_df = pd.DataFrame(report, columns=["المشكلة", "عدد الحالات", "الحل المطبق"])
with pd.ExcelWriter(OUT, engine="openpyxl", datetime_format="YYYY-MM-DD") as xw:
    df.to_excel(xw, sheet_name="البيانات النظيفة", index=False)
    rep_df.to_excel(xw, sheet_name="تقرير التنظيف", index=False)
    pd.DataFrame().to_excel(xw, sheet_name="ملخص", index=False)

wb = load_workbook(OUT)
FONT = "Arial"
head_fill = PatternFill("solid", fgColor="1F4E78")
thin = Side(style="thin", color="D0D0D0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_table(ws, start_row, ncols, nrows, start_col=1):
    for c in range(start_col, start_col + ncols):
        cell = ws.cell(row=start_row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for r in range(start_row + 1, start_row + nrows + 1):
        for c in range(start_col, start_col + ncols):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)


# شيت البيانات
ws = wb["البيانات النظيفة"]
ws.sheet_view.rightToLeft = True
style_table(ws, 1, ws.max_column, ws.max_row - 1)
widths = [11, 18, 15, 28, 14, 14, 8, 11, 14]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for r in range(2, ws.max_row + 1):
    ws.cell(r, 9).number_format = "#,##0.00"
    ws.cell(r, 3).number_format = "@"
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

# شيت التقرير
wr = wb["تقرير التنظيف"]
wr.sheet_view.rightToLeft = True
style_table(wr, 1, 3, wr.max_row - 1)
wr.column_dimensions["A"].width = 52
wr.column_dimensions["B"].width = 14
wr.column_dimensions["C"].width = 70
for r in range(2, wr.max_row + 1):
    wr.cell(r, 2).alignment = Alignment(horizontal="center", vertical="center")
    wr.cell(r, 2).number_format = "#,##0"
data_end = wr.max_row
last = data_end + 2
wr.cell(last, 1, "إجمالي الحالات المعالجة").font = Font(name=FONT, bold=True)
wr.cell(last, 2, f"=SUM(B2:B{data_end})").font = Font(name=FONT, bold=True)
wr.cell(last, 2).number_format = "#,##0"
wr.cell(last, 2).alignment = Alignment(horizontal="center")
wr.cell(last + 1, 1, "ملاحظة: البيانات تجريبية تم توليدها برمجياً لأغراض المشروع التطبيقي.").font = Font(name=FONT, italic=True, size=9, color="666666")
wr.freeze_panes = "A2"

# شيت الملخص
wsu = wb["ملخص"]
wsu.sheet_view.rightToLeft = True
wsu["A1"] = "ملخص نتائج التنظيف والتحليل"
wsu["A1"].font = Font(name=FONT, bold=True, size=14, color="1F4E78")

# قبل/بعد
wsu["A3"], wsu["B3"], wsu["C3"] = "المؤشر", "قبل", "بعد"
for i, rw in enumerate(ba.itertuples(index=False), start=4):
    wsu.cell(i, 1, rw[0]); wsu.cell(i, 2, rw[1]); wsu.cell(i, 3, rw[2])
style_table(wsu, 3, 3, 2)

# العملاء حسب المدينة
r0 = 8
for j, h in enumerate(["المدينة", "عدد العملاء"], 1):
    wsu.cell(r0, j, h)
for i, rw in enumerate(by_city.itertuples(index=False), start=r0 + 1):
    wsu.cell(i, 1, rw[0]); wsu.cell(i, 2, rw[1])
style_table(wsu, r0, 2, len(by_city))
city_end = r0 + len(by_city)

# الجنس
g0 = 8
for j, h in enumerate(["الجنس", "عدد العملاء"], 4):
    wsu.cell(g0, j, h)
for i, rw in enumerate(by_gender.itertuples(index=False), start=g0 + 1):
    wsu.cell(i, 4, rw[0]); wsu.cell(i, 5, rw[1])
style_table(wsu, g0, 2, len(by_gender), start_col=4)
g_end = g0 + len(by_gender)

# الإنفاق حسب المدينة
s0 = city_end + 3
for j, h in enumerate(["المدينة", "إجمالي الإنفاق"], 1):
    wsu.cell(s0, j, h)
for i, rw in enumerate(spent_city.itertuples(index=False), start=s0 + 1):
    wsu.cell(i, 1, rw[0]); wsu.cell(i, 2, rw[1]); wsu.cell(i, 2).number_format = "#,##0"
style_table(wsu, s0, 2, len(spent_city))
for i in range(s0 + 1, s0 + 1 + len(spent_city)):
    wsu.cell(i, 2).number_format = "#,##0"
s_end = s0 + len(spent_city)

# التسجيلات الشهرية
m0 = s_end + 3
for j, h in enumerate(["الشهر", "عدد التسجيلات"], 1):
    wsu.cell(m0, j, h)
for i, rw in enumerate(by_month.itertuples(index=False), start=m0 + 1):
    wsu.cell(i, 1, rw[0]); wsu.cell(i, 2, rw[1])
style_table(wsu, m0, 2, len(by_month))
m_end = m0 + len(by_month)

wsu.column_dimensions["A"].width = 28
wsu.column_dimensions["B"].width = 16
wsu.column_dimensions["C"].width = 14
wsu.column_dimensions["D"].width = 14
wsu.column_dimensions["E"].width = 14

# رسم: العملاء حسب المدينة
ch1 = BarChart(); ch1.type = "col"; ch1.title = "عدد العملاء حسب المدينة"
ch1.add_data(Reference(wsu, min_col=2, min_row=r0, max_row=city_end), titles_from_data=True)
ch1.set_categories(Reference(wsu, min_col=1, min_row=r0 + 1, max_row=city_end))
ch1.legend = None; ch1.height = 8; ch1.width = 16
ch1.y_axis.delete = False; ch1.x_axis.delete = False
ch1.dataLabels = DataLabelList(); ch1.dataLabels.showVal = True
ch1.dataLabels.showSerName = False; ch1.dataLabels.showCatName = False; ch1.dataLabels.showLegendKey = False
wsu.add_chart(ch1, "G3")

# رسم: الجنس
ch2 = PieChart(); ch2.title = "توزيع العملاء حسب الجنس"
ch2.add_data(Reference(wsu, min_col=5, min_row=g0, max_row=g_end), titles_from_data=True)
ch2.set_categories(Reference(wsu, min_col=4, min_row=g0 + 1, max_row=g_end))
ch2.dataLabels = DataLabelList(); ch2.dataLabels.showPercent = True
ch2.dataLabels.showVal = False; ch2.dataLabels.showSerName = False; ch2.dataLabels.showCatName = False; ch2.dataLabels.showLegendKey = False
ch2.height = 8; ch2.width = 10
wsu.add_chart(ch2, "G20")

# رسم: التسجيلات الشهرية
ch3 = LineChart(); ch3.title = "التسجيلات الشهرية"
ch3.add_data(Reference(wsu, min_col=2, min_row=m0, max_row=m_end), titles_from_data=True)
ch3.set_categories(Reference(wsu, min_col=1, min_row=m0 + 1, max_row=m_end))
ch3.legend = None; ch3.height = 8; ch3.width = 16
for _s in ch3.series: _s.smooth = False
ch3.y_axis.delete = False; ch3.x_axis.delete = False
wsu.add_chart(ch3, "G37")

# رسم: الإنفاق حسب المدينة
ch4 = BarChart(); ch4.type = "bar"; ch4.title = "إجمالي الإنفاق حسب المدينة"
ch4.add_data(Reference(wsu, min_col=2, min_row=s0, max_row=s_end), titles_from_data=True)
ch4.set_categories(Reference(wsu, min_col=1, min_row=s0 + 1, max_row=s_end))
ch4.legend = None; ch4.height = 8; ch4.width = 16
ch4.y_axis.delete = False; ch4.x_axis.delete = False
wsu.add_chart(ch4, "G54")

wb.move_sheet("ملخص", offset=-2)
wb.active = 0
wb.save(OUT)

print(f"قبل: {rows_before} صف | بعد: {rows_after} صف")
print(rep_df.to_string(index=False))
